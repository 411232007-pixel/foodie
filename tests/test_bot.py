"""The Discord layer, tested with stand-in message objects (no network, no real Discord).
These tests show the program's own logic is right; they do NOT replace trying it in a real server."""

import asyncio
import threading
import time
from types import SimpleNamespace
from unittest.mock import AsyncMock

import discord
import pytest

import foodie.bot as bot_module
from foodie.bot import FoodieBot
from foodie.config import MAX_IMAGES_PER_MESSAGE, MAX_QUEUE_JOBS, RATE_LIMIT_PHOTOS
from foodie.processor import PhotoResult
from tests.samples import PNG

CHANNEL = 123
OK = PhotoResult("ok", is_food=True, name="牛肉麵", ingredients=("牛肉",), weight_g=500, calories_kcal=650, daily_total=650)


def attachment(name="a.png", size=100, content_type="image/png"):
    return SimpleNamespace(filename=name, size=size, content_type=content_type, read=AsyncMock(return_value=PNG))


def message(attachments, channel=CHANNEL, author_id=7, is_bot=False, content="", label="m"):
    return SimpleNamespace(
        label=label,
        author=SimpleNamespace(bot=is_bot, id=author_id, display_name="小明"),
        channel=SimpleNamespace(id=channel),
        attachments=attachments,
        content=content,
        reply=AsyncMock(),
        add_reaction=AsyncMock(),
        remove_reaction=AsyncMock(),
    )


@pytest.fixture
def bot(cfg, ledger):
    instance = FoodieBot(cfg, ledger)
    instance._connection.user = SimpleNamespace(id=999, name="Foodie")
    return instance


def run_worker(bot, coro_factory):
    """Start the queue worker, run the scenario, wait until the queue is empty, stop the worker."""

    async def scenario():
        await bot.setup_hook()
        await coro_factory()
        await asyncio.wait_for(bot.queue.join(), timeout=10)
        bot._worker.cancel()

    asyncio.run(scenario())


# ---- which messages are looked at ------------------------------------------------------------

@pytest.mark.parametrize(
    "msg",
    [
        message([attachment()], channel=456),  # another channel
        message([attachment()], is_bot=True),  # a bot (including ourselves)
        message([]),  # plain text
        message([attachment("notes.pdf", content_type="application/pdf")]),
        message([attachment("clip.mp4", content_type="video/mp4")]),
    ],
)
def test_only_images_in_the_chosen_channel_are_handled(bot, msg):
    asyncio.run(bot.on_message(msg))
    assert bot.queue.empty()
    msg.add_reaction.assert_not_called()
    msg.reply.assert_not_called()


def test_an_image_gets_hourglass_then_one_reply_then_hourglass_removed(bot, monkeypatch):
    monkeypatch.setattr(bot_module, "process_image", lambda *a, **k: OK)
    msg = message([attachment()])
    run_worker(bot, lambda: bot.on_message(msg))
    msg.add_reaction.assert_awaited_once_with("⏳")
    msg.remove_reaction.assert_awaited_once()
    assert msg.remove_reaction.await_args.args[0] == "⏳"
    msg.reply.assert_awaited_once()
    assert "你今天已攝取熱量:約 650 大卡" in msg.reply.await_args.args[0]


def test_reply_pings_only_the_author(bot, monkeypatch):
    monkeypatch.setattr(bot_module, "process_image", lambda *a, **k: OK)
    msg = message([attachment()])
    run_worker(bot, lambda: bot.on_message(msg))
    mentions = msg.reply.await_args.kwargs["allowed_mentions"]
    assert (mentions.everyone, mentions.roles, mentions.users, mentions.replied_user) == (False, False, False, True)


# ---- one claude at a time ----------------------------------------------------------------

def test_photos_are_analysed_strictly_one_at_a_time_across_messages(bot, monkeypatch):
    state = {"now": 0, "max": 0, "order": []}
    lock = threading.Lock()

    def slow(cfg, ledger, data, uploader, uploader_id, caption, *a, **k):
        with lock:
            state["now"] += 1
            state["max"] = max(state["max"], state["now"])
            state["order"].append(caption)
        time.sleep(0.03)
        with lock:
            state["now"] -= 1
        return OK

    monkeypatch.setattr(bot_module, "process_image", slow)
    first = message([attachment(), attachment(), attachment()], content="m1", author_id=1)
    second = message([attachment(), attachment()], content="m2", author_id=2)
    third = message([attachment()], content="m3", author_id=3)

    async def send_all():
        for m in (first, second, third):
            await bot.on_message(m)

    run_worker(bot, send_all)
    assert state["max"] == 1  # never two analyses at the same time
    assert state["order"] == ["m1"] * 3 + ["m2"] * 2 + ["m3"]  # first come, first served, photo by photo
    for m in (first, second, third):
        m.reply.assert_awaited_once()  # exactly one combined reply per message
    assert "3 張照片" in first.reply.await_args.args[0]


def test_the_discord_user_id_not_the_display_name_is_what_identifies_the_uploader(bot, monkeypatch):
    seen = []
    monkeypatch.setattr(bot_module, "process_image", lambda cfg, ledger, data, name, uid, cap, *a, **k: seen.append((name, uid)) or OK)
    msg = message([attachment()], author_id=1562019939534110794)
    run_worker(bot, lambda: bot.on_message(msg))
    assert seen == [("小明", "1562019939534110794")]


# ---- failures never stop the queue --------------------------------------------------------

def test_a_crashing_photo_gets_the_failure_text_and_later_messages_still_work(bot, monkeypatch):
    def flaky(cfg, ledger, data, uploader, uploader_id, caption, *a, **k):
        if caption == "boom":
            raise RuntimeError("unexpected")
        return OK

    monkeypatch.setattr(bot_module, "process_image", flaky)
    bad = message([attachment()], content="boom")
    good = message([attachment()], content="fine")

    async def send():
        await bot.on_message(bad)
        await bot.on_message(good)

    run_worker(bot, send)
    assert bad.reply.await_args.args[0] == "分析失敗，請再傳一次"
    assert "是食物:牛肉麵" in good.reply.await_args.args[0]
    for m in (bad, good):
        m.remove_reaction.assert_awaited_once()  # no message is left with ⏳


def test_an_analysis_failure_is_a_reply_to_the_original_message(bot, monkeypatch):
    monkeypatch.setattr(bot_module, "process_image", lambda *a, **k: PhotoResult("failed"))
    msg = message([attachment()])
    run_worker(bot, lambda: bot.on_message(msg))
    assert msg.reply.await_args.args[0] == "分析失敗，請再傳一次"


def test_discord_errors_while_replying_or_reacting_do_not_kill_the_worker(bot, monkeypatch):
    monkeypatch.setattr(bot_module, "process_image", lambda *a, **k: OK)
    broken = message([attachment()])
    broken.reply = AsyncMock(side_effect=RuntimeError("network down"))
    broken.remove_reaction = AsyncMock(side_effect=RuntimeError("message deleted"))
    fine = message([attachment()])

    async def send():
        await bot.on_message(broken)
        await bot.on_message(fine)

    run_worker(bot, send)
    fine.reply.assert_awaited_once()


def test_missing_add_reactions_permission_is_survivable(bot, monkeypatch, caplog):
    monkeypatch.setattr(bot_module, "process_image", lambda *a, **k: OK)
    forbidden = discord.Forbidden(SimpleNamespace(status=403, reason="Forbidden"), "Missing Permissions")
    msgs = [message([attachment()]) for _ in range(2)]
    for m in msgs:
        m.add_reaction = AsyncMock(side_effect=forbidden)

    async def send():
        for m in msgs:
            await bot.on_message(m)

    run_worker(bot, send)
    for m in msgs:
        m.reply.assert_awaited_once()  # still answered
    assert caplog.text.count("沒有 Add Reactions 權限") == 1  # warned once, not for every photo


# ---- !export -------------------------------------------------------------------------------

def fill_log(bot, uploader_id="7", uploader="小明"):
    from datetime import datetime

    from foodie.ledger import LogRecord

    for title, kcal in (("牛肉麵", 650.0), ("紅茶", 120.0)):
        bot.ledger.append(LogRecord(datetime(2026, 10, 20, 12, 0, 0), uploader, uploader_id, title, "x", 100.0, kcal, "photos\\a.jpg"))
    bot.ledger.append(LogRecord(datetime(2026, 10, 20, 12, 5, 0), "小華", "8", "別人的飯", "x", 100.0, 300.0, "photos\\b.jpg"))


def test_export_sends_the_askers_own_history_as_an_xlsx_reply(bot):
    from openpyxl import load_workbook

    fill_log(bot)
    msg = message([], content="!export", author_id=7)
    asyncio.run(bot.on_message(msg))
    msg.reply.assert_awaited_once()
    kwargs = msg.reply.await_args.kwargs
    assert kwargs["file"].filename.startswith("foodie-") and kwargs["file"].filename.endswith(".xlsx")
    assert "共 2 筆" in msg.reply.await_args.args[0]
    sheet = load_workbook(kwargs["file"].fp).active
    titles = [r[2] for r in sheet.iter_rows(min_row=2, values_only=True)]
    assert titles == ["牛肉麵", "紅茶"]  # not 別人的飯
    assert (kwargs["allowed_mentions"].everyone, kwargs["allowed_mentions"].replied_user) == (False, True)
    msg.add_reaction.assert_not_called()  # no hourglass, no queue
    assert bot.queue.empty()


@pytest.mark.parametrize("text", ["!EXPORT", "  !export  ", "!Export\n"])
def test_the_command_ignores_case_and_surrounding_space(bot, text):
    fill_log(bot)
    msg = message([], content=text, author_id=7)
    asyncio.run(bot.on_message(msg))
    assert "file" in msg.reply.await_args.kwargs


@pytest.mark.parametrize(
    "msg",
    [
        message([], content="!export", channel=456),  # another channel
        message([], content="!export", is_bot=True),  # a bot
        message([], content="!export please"),  # not exactly the command
        message([], content="export"),
        message([], content="!exports"),
    ],
)
def test_other_texts_do_not_trigger_an_export(bot, msg):
    fill_log(bot)
    asyncio.run(bot.on_message(msg))
    msg.reply.assert_not_called()


def test_export_with_no_records_says_so(bot):
    msg = message([], content="!export", author_id=7)
    asyncio.run(bot.on_message(msg))
    assert msg.reply.await_args.args[0] == "你還沒有任何紀錄，先傳一張食物照片吧"
    assert "file" not in msg.reply.await_args.kwargs


def test_exports_are_rate_limited_per_person(bot):
    fill_log(bot)

    async def scenario():
        first = message([], content="!export", author_id=7)
        second = message([], content="!export", author_id=7)
        other = message([], content="!export", author_id=8)
        for m in (first, second, other):
            await bot.on_message(m)
        return first, second, other

    first, second, other = asyncio.run(scenario())
    assert "file" in first.reply.await_args.kwargs
    assert second.reply.await_args.args[0] == "匯出太頻繁了，請等一分鐘再試" and "file" not in second.reply.await_args.kwargs
    assert "file" in other.reply.await_args.kwargs  # someone else is not affected


def test_missing_attach_files_permission_falls_back_to_a_plain_message(bot, caplog):
    fill_log(bot)
    forbidden = discord.Forbidden(SimpleNamespace(status=403, reason="Forbidden"), "Missing Permissions")
    msg = message([], content="!export", author_id=7)
    msg.reply = AsyncMock(side_effect=[forbidden, None])
    asyncio.run(bot.on_message(msg))
    assert msg.reply.await_count == 2
    assert "Attach Files" in msg.reply.await_args.args[0] and "file" not in msg.reply.await_args.kwargs
    assert "Attach Files" in caplog.text


def test_an_unreadable_log_gives_a_failure_message_not_a_crash(bot, monkeypatch):
    monkeypatch.setattr(bot.ledger, "export_for", lambda *a: (_ for _ in ()).throw(RuntimeError("broken")))
    msg = message([], content="!export", author_id=7)
    asyncio.run(bot.on_message(msg))
    assert msg.reply.await_args.args[0] == "匯出失敗，請稍後再試"


def test_a_photo_with_the_command_as_its_caption_is_still_a_photo(bot, monkeypatch):
    monkeypatch.setattr(bot_module, "process_image", lambda *a, **k: OK)
    msg = message([attachment()], content="!export")
    run_worker(bot, lambda: bot.on_message(msg))
    msg.add_reaction.assert_awaited_once_with("⏳")
    assert "是食物:牛肉麵" in msg.reply.await_args.args[0]


# ---- limits -------------------------------------------------------------------------------

def test_only_the_first_four_photos_of_a_message_are_processed(bot, monkeypatch):
    calls = []
    monkeypatch.setattr(bot_module, "process_image", lambda *a, **k: calls.append(1) or OK)
    msg = message([attachment() for _ in range(MAX_IMAGES_PER_MESSAGE + 2)])
    run_worker(bot, lambda: bot.on_message(msg))
    assert len(calls) == MAX_IMAGES_PER_MESSAGE
    assert "沒有處理" in msg.reply.await_args.args[0]


def test_full_queue_is_refused_politely_without_hourglass(bot):
    async def scenario():
        for i in range(MAX_QUEUE_JOBS):
            await bot.on_message(message([attachment()], author_id=1000 + i))  # no worker running: they pile up
        extra = message([attachment()], author_id=5000)
        await bot.on_message(extra)
        return extra

    extra = asyncio.run(scenario())
    assert bot.queue.qsize() == MAX_QUEUE_JOBS
    extra.reply.assert_awaited_once()
    assert "太忙" in extra.reply.await_args.args[0]
    extra.add_reaction.assert_not_called()


def test_one_person_cannot_flood(bot):
    async def scenario():
        sent = []
        for _ in range(RATE_LIMIT_PHOTOS + 1):
            m = message([attachment()], author_id=42)
            sent.append(m)
            await bot.on_message(m)
        return sent

    sent = asyncio.run(scenario())
    assert bot.queue.qsize() == RATE_LIMIT_PHOTOS
    assert "頻繁" in sent[-1].reply.await_args.args[0]
    sent[-1].add_reaction.assert_not_called()


def test_oversized_image_is_refused_without_downloading(bot, monkeypatch):
    called = []
    monkeypatch.setattr(bot_module, "process_image", lambda *a, **k: called.append(1) or OK)
    big = attachment(size=bot_module.MAX_IMAGE_BYTES + 1)
    msg = message([big])
    run_worker(bot, lambda: bot.on_message(msg))
    big.read.assert_not_called()
    assert not called and "太大" in msg.reply.await_args.args[0]


# ---- shutting down -----------------------------------------------------------------------

def test_shutdown_takes_the_hourglass_off_everything_still_waiting(bot):
    waiting = [message([attachment()], author_id=i) for i in range(3)]

    async def scenario():
        for m in waiting:
            await bot.on_message(m)
        await bot.setup_hook()  # worker exists but has not run yet
        await bot.close()

    asyncio.run(scenario())
    for m in waiting:
        m.remove_reaction.assert_awaited_once()
