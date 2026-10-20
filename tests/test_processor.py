from datetime import datetime

import pytest
from openpyxl import load_workbook

from foodie.ledger import LedgerBusy
from foodie.processor import process_image
from foodie.replies import build_reply
from foodie.result_check import AnalysisError, FoodResult, NotLegible
from tests.samples import JPEG

NOW = datetime(2026, 10, 20, 12, 0, 0)


def food(**kw) -> FoodResult:
    base = dict(is_food=True, legible=True, name="牛肉麵", ingredients=("牛肉", "麵條"), weight_g=500.0, calories_kcal=650.0)
    return FoodResult(**{**base, **kw})


def run(cfg, ledger, result=None, error=None, data=JPEG, uploader="小明"):
    def fake_analyze(_cfg, _data, _ext, _caption):
        if error:
            raise error
        return result

    return process_image(cfg, ledger, data, uploader, "", now=NOW, analyze_fn=fake_analyze)


def sheet_rows(cfg):
    return [list(r) for r in load_workbook(cfg.ledger_path).active.iter_rows(min_row=2, values_only=True)]


def test_food_photo_is_recorded_and_saved(cfg, ledger):
    result = run(cfg, ledger, food())
    assert result.status == "ok" and result.daily_total == 650
    (row,) = sheet_rows(cfg)
    assert row[2:7] == ["牛肉麵", "牛肉、麵條", 500.0, 650.0, 650.0]
    saved = cfg.data_dir / row[7]
    assert saved.name == "2026-10-20 12-00-00 牛肉麵.jpg" and saved.read_bytes() == JPEG


def test_not_food_is_recorded_too_with_zero_numbers(cfg, ledger):
    result = run(cfg, ledger, food(is_food=False, name="馬克杯", ingredients=(), weight_g=300.0, calories_kcal=999.0))
    assert result.status == "ok" and result.daily_total == 0
    (row,) = sheet_rows(cfg)
    assert row[2:7] == ["不是食物", "馬克杯", 0.0, 0.0, 0.0]
    assert (cfg.data_dir / row[7]).name == "2026-10-20 12-00-00 不是食物.jpg"
    reply = build_reply([result])
    for part in ("你傳的不是食物,是馬克杯", "成份:馬克杯", "重量:0 公克", "熱量:0 大卡", "今日已攝取熱量(全頻道累計):約 0 大卡"):
        assert part in reply


@pytest.mark.parametrize("error", [AnalysisError("bad"), RuntimeError("surprise")])
def test_analysis_trouble_means_failure_and_nothing_is_kept(cfg, ledger, error):
    assert run(cfg, ledger, error=error).status == "failed"
    assert not cfg.ledger_path.exists()
    assert not list(cfg.photos_dir.glob("*")) if cfg.photos_dir.exists() else True


def test_blurry_photo_is_not_kept(cfg, ledger):
    assert run(cfg, ledger, error=NotLegible("blurry")).status == "unclear"
    assert not cfg.ledger_path.exists()


def test_not_an_image_is_refused_without_calling_claude(cfg, ledger):
    assert run(cfg, ledger, error=AssertionError("must not be called"), data=b"MZ not an image").status == "failed"


def test_not_food_does_not_change_the_days_total(cfg, ledger):
    run(cfg, ledger, food(calories_kcal=500.0))
    after = run(cfg, ledger, food(is_food=False, name="貓", ingredients=(), weight_g=4000.0, calories_kcal=4000.0))
    assert after.daily_total == 500
    assert run(cfg, ledger, food(calories_kcal=100.0)).daily_total == 600


def test_extension_follows_the_real_content_never_just_renamed(cfg, ledger):
    from tests.samples import PNG

    run(cfg, ledger, food(name="一"), data=JPEG)
    run(cfg, ledger, food(name="二"), data=PNG)
    names = sorted(p.name for p in cfg.photos_dir.glob("*"))
    assert names == ["2026-10-20 12-00-00 一.jpg", "2026-10-20 12-00-00 二.png"]
    for photo in cfg.photos_dir.glob("*"):
        assert photo.read_bytes()[:3] == (b"\xff\xd8\xff" if photo.suffix == ".jpg" else b"\x89PN")


def test_if_the_excel_cannot_be_written_the_photo_is_removed_too(cfg, ledger, monkeypatch):
    cfg.photos_dir.mkdir(parents=True)

    def busy(_record):
        raise LedgerBusy("open in Excel")

    monkeypatch.setattr(ledger, "append", busy)
    assert run(cfg, ledger, food()).status == "ledger_busy"
    assert list(cfg.photos_dir.glob("*")) == []


def test_two_photos_in_a_row_build_on_each_others_total(cfg, ledger):
    first = run(cfg, ledger, food(calories_kcal=600.0), uploader="小明")
    second = run(cfg, ledger, food(name="紅茶", calories_kcal=120.0), uploader="小華")
    assert (first.daily_total, second.daily_total) == (600, 720)
    reply = build_reply([first, second])
    assert "這則訊息的 2 張照片" in reply and "今日已攝取熱量(全頻道累計):約 720 大卡" in reply
    assert len(list(cfg.photos_dir.glob("*"))) == 2


def test_evil_food_name_cannot_escape_or_become_a_formula(cfg, ledger):
    run(cfg, ledger, food(name="=cmd|' /C calc'!A0 ..\\..\\evil"))
    (row,) = sheet_rows(cfg)
    assert (cfg.data_dir / row[7]).parent == cfg.photos_dir
    cell = load_workbook(cfg.ledger_path).active.cell(row=2, column=3)
    assert cell.data_type == "s"


def test_single_food_reply_has_everything(cfg, ledger):
    text = build_reply([run(cfg, ledger, food())])
    for part in ("是食物:牛肉麵", "牛肉、麵條", "約 500 公克", "約 650 大卡", "今日已攝取熱量(全頻道累計):約 650 大卡"):
        assert part in text


def test_failure_replies(cfg, ledger):
    assert build_reply([run(cfg, ledger, error=AnalysisError("x"))]) == "分析失敗，請再傳一次"
    assert "看不清楚" in build_reply([run(cfg, ledger, error=NotLegible("x"))])


def test_mentions_in_names_are_defused(cfg, ledger):
    text = build_reply([run(cfg, ledger, food(name="@everyone 來吃"))])
    assert "@everyone" not in text.replace("@​everyone", "")
