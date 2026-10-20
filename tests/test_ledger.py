import zipfile
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

import foodie.ledger as ledger_module
from foodie.ledger import HEADERS, Ledger, LedgerBusy, LogRecord


def record(when="2026-10-20 12:00:00", uploader="小明", uploader_id="1001", title="牛肉麵", kcal=650.0, **kw) -> LogRecord:
    defaults = dict(
        when=datetime.fromisoformat(when),
        uploader=uploader,
        uploader_id=uploader_id,
        title=title,
        ingredients="牛肉、麵條",
        weight_g=500.0,
        calories_kcal=kcal,
        photo_path="photos\\x.jpg",
    )
    return LogRecord(**{**defaults, **kw})


def rows(ledger):
    sheet = load_workbook(ledger.path).active
    return [list(r) for r in sheet.iter_rows(values_only=True)]


def test_first_write_creates_file_with_headers(ledger):
    ledger.append(record())
    data = rows(ledger)
    assert data[0] == HEADERS
    assert data[1][:8] == ["2026-10-20 12:00:00", "小明", "牛肉麵", "牛肉、麵條", 500.0, 650.0, 650.0, "photos\\x.jpg"]
    assert data[1][8] == "1001"  # the id sits in a 9th column so the 8 required columns keep their places


def test_rows_are_added_at_the_bottom_never_overwritten(ledger):
    ledger.append(record(title="一"))
    ledger.append(record(title="二"))
    ledger2 = Ledger(ledger.path)  # a fresh program start must keep appending
    ledger2.append(record(title="三"))
    assert [r[2] for r in rows(ledger)[1:]] == ["一", "二", "三"]


def test_daily_total_is_kept_per_person(ledger):
    assert ledger.append(record(uploader="小明", uploader_id="1", kcal=600)) == 600
    assert ledger.append(record(uploader="小華", uploader_id="2", kcal=250.5)) == 250.5  # not 850.5
    assert ledger.append(record(uploader="小明", uploader_id="1", kcal=100)) == 700
    assert ledger.append(record(uploader="小明", uploader_id="1", kcal=0, title="不是食物")) == 700
    assert ledger.append(record(uploader="小華", uploader_id="2", kcal=50)) == 300.5
    assert [r[6] for r in rows(ledger)[1:]] == [600, 250.5, 700, 700, 300.5]


def test_two_people_with_the_same_display_name_are_not_mixed_up(ledger):
    ledger.append(record(uploader="阿明", uploader_id="1", kcal=600))
    assert ledger.append(record(uploader="阿明", uploader_id="2", kcal=200)) == 200


def test_renaming_does_not_lose_the_days_total(ledger):
    ledger.append(record(uploader="舊暱稱", uploader_id="1", kcal=600))
    assert ledger.append(record(uploader="新暱稱", uploader_id="1", kcal=100)) == 700


def test_rows_from_before_the_id_column_are_matched_by_name(ledger):
    # A log made by the earlier version: eight columns, no id, header without 上傳者ID.
    old = Workbook()
    sheet = old.active
    sheet.append(HEADERS[:8])
    sheet.append(["2026-10-20 08:00:00", "小明", "早餐", "", 100.0, 300.0, 300.0, "photos\\a.jpg"])
    sheet.append(["2026-10-20 09:00:00", "小華", "早餐", "", 100.0, 500.0, 800.0, "photos\\b.jpg"])
    old.save(ledger.path)
    assert ledger.append(record(uploader="小明", uploader_id="1", kcal=100)) == 400
    assert ledger.append(record(uploader="小華", uploader_id="2", kcal=10)) == 510
    data = rows(ledger)
    assert data[0] == HEADERS  # the missing header was added
    assert [r[8] for r in data[1:]] == [None, None, "1", "2"]
    assert data[1][:8] == ["2026-10-20 08:00:00", "小明", "早餐", None, 100.0, 300.0, 300.0, "photos\\a.jpg"]  # untouched


def test_a_19_digit_id_is_stored_without_losing_digits(ledger):
    ledger.append(record(uploader_id="1562019939534110794"))
    assert rows(ledger)[1][8] == "1562019939534110794"


def test_daily_total_restarts_at_midnight(ledger):
    ledger.append(record(when="2026-10-20 23:59:59", kcal=700))
    assert ledger.append(record(when="2026-10-21 00:00:00", kcal=300)) == 300
    assert ledger.append(record(when="2026-10-21 08:00:00", kcal=200)) == 500


def test_a_row_edited_by_hand_is_used_by_later_totals(ledger):
    ledger.append(record(kcal=600))
    workbook = load_workbook(ledger.path)
    workbook.active.cell(row=2, column=6, value=100)
    workbook.save(ledger.path)
    assert ledger.append(record(kcal=50)) == 150


def test_a_date_cell_that_excel_turned_into_a_real_date_still_counts(ledger):
    ledger.append(record(kcal=600))
    workbook = load_workbook(ledger.path)
    workbook.active.cell(row=2, column=1, value=datetime(2026, 10, 20, 12, 0, 0))  # what retyping the cell does
    workbook.save(ledger.path)
    assert ledger.append(record(kcal=50)) == 650


def test_no_temp_file_is_left_when_saving_fails(ledger, monkeypatch):
    ledger.append(record())

    def boom(self, filename):
        Path(filename).write_bytes(b"half written")
        raise PermissionError("locked")

    monkeypatch.setattr(ledger_module.Workbook, "save", boom)
    with pytest.raises(LedgerBusy):
        ledger.append(record())
    assert not ledger.path.with_name(ledger.path.name + ".tmp").exists()


@pytest.mark.parametrize("evil", ["=1+1", '=HYPERLINK("http://evil.example","click")', "+cmd", "-2", "@SUM(1)", "\t=1"])
def test_formula_like_text_is_stored_as_plain_text(ledger, evil):
    ledger.append(record(uploader=evil, title=evil, ingredients=evil))
    sheet = load_workbook(ledger.path).active
    for column in (2, 3, 4):
        cell = sheet.cell(row=2, column=column)
        assert cell.data_type == "s", f"column {column} became {cell.data_type}"
        assert cell.value == evil
        assert cell.quotePrefix
    with zipfile.ZipFile(ledger.path) as archive:
        assert b"<f>" not in archive.read("xl/worksheets/sheet1.xml")  # no formula element anywhere


def test_ordinary_text_gets_no_quote_prefix(ledger):
    ledger.append(record())
    assert not load_workbook(ledger.path).active.cell(row=2, column=3).quotePrefix


def test_locked_file_is_retried_then_reported(ledger, monkeypatch, caplog):
    ledger.append(record())
    calls = []

    def locked(src, dst):
        calls.append(1)
        raise PermissionError("file is open in Excel")

    monkeypatch.setattr(ledger_module.os, "replace", locked)
    with pytest.raises(LedgerBusy):
        ledger.append(record(title="新的"))
    assert len(calls) == 3
    assert "Excel" in caplog.text
    assert [r[2] for r in rows(ledger)[1:]] == ["牛肉麵"]  # old data untouched, nothing half-written
    assert not ledger.path.with_name(ledger.path.name + ".tmp").exists()


def test_a_lock_that_clears_on_retry_succeeds(ledger, monkeypatch):
    ledger.append(record())
    real_replace, attempts = ledger_module.os.replace, []

    def flaky(src, dst):
        attempts.append(1)
        if len(attempts) < 3:
            raise PermissionError("busy")
        return real_replace(src, dst)

    monkeypatch.setattr(ledger_module.os, "replace", flaky)
    assert ledger.append(record(kcal=100)) == 750
    assert len(rows(ledger)) == 3


def test_a_broken_file_is_never_overwritten(ledger):
    ledger.path.write_bytes(b"this is not an xlsx file")
    with pytest.raises(Exception):
        ledger.append(record())
    assert ledger.path.read_bytes() == b"this is not an xlsx file"
