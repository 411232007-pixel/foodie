"""The Excel log. Holds 防護 4 (see docs/architecture.md section 12)."""

from __future__ import annotations

import contextlib
import io
import logging
import os
import threading
import time
from collections.abc import Callable
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from foodie.config import LEDGER_RETRIES, LEDGER_RETRY_SECONDS

log = logging.getLogger("foodie.ledger")

HEADERS = [
    "日期時間",
    "上傳者",
    "食物敘述",
    "成份",
    "重量(公克)",
    "熱量(大卡)",
    "當日已攝取熱量(大卡)",
    "照片存放位置",
    "上傳者ID",  # last on purpose: the eight required columns keep their positions
]
COL_WHEN, COL_UPLOADER, COL_TITLE, COL_INGREDIENTS, COL_WEIGHT, COL_CALORIES, COL_DAILY, COL_PHOTO, COL_UPLOADER_ID = range(1, 10)
COLUMN_WIDTHS = [20, 16, 22, 36, 12, 12, 22, 60, 22]
EXPORT_HEADERS = HEADERS[:7]  # no photo path (a path on this computer) and no id
EXPORT_NUMBER_COLUMNS = {COL_WEIGHT, COL_CALORIES, COL_DAILY}
FORMULA_STARTERS = ("=", "+", "-", "@", "\t", "\r", "\n")


class LedgerBusy(Exception):
    """The file could not be written, usually because it is open in Excel."""


@dataclass(frozen=True)
class LogRecord:
    when: datetime
    uploader: str  # display name: shown to people, can change or repeat
    uploader_id: str  # Discord user id: never changes, so the daily total is kept per person by this
    title: str
    ingredients: str
    weight_g: float
    calories_kcal: float
    photo_path: str


class Ledger:
    def __init__(self, path: Path, sleep: Callable[[float], None] = time.sleep) -> None:
        self.path = path
        self._sleep = sleep
        self._lock = threading.Lock()

    def append(self, record: LogRecord) -> float:
        """Add one row at the bottom. Returns today's total for THIS uploader only."""
        with self._lock:
            for attempt in range(1, LEDGER_RETRIES + 1):
                try:
                    return self._append_once(record)
                except PermissionError as exc:
                    log.warning(
                        "⚠️ 寫不進 %s(第 %d/%d 次)。是不是正開著 Excel?請把它關掉,%d 秒後重試。",
                        self.path.name, attempt, LEDGER_RETRIES, LEDGER_RETRY_SECONDS,
                    )
                    if attempt == LEDGER_RETRIES:
                        raise LedgerBusy(str(exc)) from exc
                    self._sleep(LEDGER_RETRY_SECONDS)
        raise AssertionError("unreachable")

    def _append_once(self, record: LogRecord) -> float:
        if self.path.exists():
            workbook = load_workbook(self.path)  # an unreadable file raises: we never start over on top of it
            sheet = workbook.active
            if sheet.cell(row=1, column=COL_UPLOADER_ID).value is None:  # a log from before the ID column existed
                sheet.cell(row=1, column=COL_UPLOADER_ID, value=HEADERS[COL_UPLOADER_ID - 1])
                sheet.column_dimensions[chr(64 + COL_UPLOADER_ID)].width = COLUMN_WIDTHS[COL_UPLOADER_ID - 1]
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "紀錄"
            sheet.append(HEADERS)
            for index, width in enumerate(COLUMN_WIDTHS, start=1):
                sheet.column_dimensions[chr(64 + index)].width = width

        today = record.when.strftime("%Y-%m-%d")
        daily_total = round(self._calories_on(sheet, today, record.uploader_id, record.uploader) + record.calories_kcal, 1)

        row = sheet.max_row + 1
        _put_text(sheet, row, COL_WHEN, record.when.strftime("%Y-%m-%d %H:%M:%S"))
        _put_text(sheet, row, COL_UPLOADER, record.uploader)
        _put_text(sheet, row, COL_TITLE, record.title)
        _put_text(sheet, row, COL_INGREDIENTS, record.ingredients)
        sheet.cell(row=row, column=COL_WEIGHT, value=round(record.weight_g, 1))
        sheet.cell(row=row, column=COL_CALORIES, value=round(record.calories_kcal, 1))
        sheet.cell(row=row, column=COL_DAILY, value=daily_total)
        _put_text(sheet, row, COL_PHOTO, record.photo_path)
        _put_text(sheet, row, COL_UPLOADER_ID, record.uploader_id)  # as text: a 19-digit id would lose digits as a number

        self._save_atomic(workbook)
        return daily_total

    @staticmethod
    def _calories_on(sheet: Worksheet, day: str, uploader_id: str, uploader: str) -> float:
        """Today's calories for one person. Rows are matched by Discord id; rows written before the
        id column existed have no id, so those fall back to matching the display name."""
        total = 0.0
        for row in sheet.iter_rows(min_row=2, values_only=True):
            when, calories = row[COL_WHEN - 1], row[COL_CALORIES - 1]
            if isinstance(when, datetime):  # Excel turns a retyped date cell into a real date
                when = when.strftime("%Y-%m-%d %H:%M:%S")
            if not (isinstance(when, str) and when.startswith(day) and _is_number(calories)):
                continue
            if _is_mine(row, uploader_id, uploader):
                total += calories
        return total

    def export_for(self, uploader_id: str, uploader: str) -> tuple[bytes, int] | None:
        """Everything one person has logged, as a brand-new .xlsx (bytes) plus the row count.

        Only rows carrying that person's Discord id go in (old rows without an id are never exported,
        see _is_mine), without the photo path (it is a path on this computer) and without the id column.
        Returns None when the person has no rows. The real log is never touched.
        """
        with self._lock:
            if not self.path.exists():
                return None
            try:
                sheet = load_workbook(self.path).active
                mine = [
                    row
                    for row in sheet.iter_rows(min_row=2, values_only=True)
                    if uploader_id and _is_mine(row, uploader_id, uploader, name_fallback=False)
                ]
            except PermissionError as exc:
                raise LedgerBusy(str(exc)) from exc
        if not mine:
            return None

        out = Workbook()
        target = out.active
        target.title = "我的紀錄"
        target.append(EXPORT_HEADERS)
        for index, width in enumerate(COLUMN_WIDTHS[: len(EXPORT_HEADERS)], start=1):
            target.column_dimensions[chr(64 + index)].width = width
        for source in mine:
            row = target.max_row + 1
            for column in range(1, len(EXPORT_HEADERS) + 1):
                value = source[column - 1] if len(source) >= column else None
                if column in EXPORT_NUMBER_COLUMNS and _is_number(value):
                    target.cell(row=row, column=column, value=value)
                else:  # every other cell, including a hand-typed "number" that is really text, goes through 防護 4
                    _put_text(target, row, column, "" if value is None else str(value))
        buffer = io.BytesIO()
        out.save(buffer)
        return buffer.getvalue(), len(mine)

    def _save_atomic(self, workbook: Workbook) -> None:
        """Save next to the real file first, then swap: a crash mid-save cannot damage old records."""
        self.path.parent.mkdir(parents=True, exist_ok=True)
        temp = self.path.with_name(self.path.name + ".tmp")
        try:
            workbook.save(temp)
            os.replace(temp, self.path)
        except BaseException:
            with contextlib.suppress(OSError):
                temp.unlink(missing_ok=True)  # never leave a half-written temp file behind
            raise


def _is_number(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _is_mine(row: tuple, uploader_id: str, uploader: str, *, name_fallback: bool = True) -> bool:
    """Does this log row belong to this person? By Discord id.

    Rows written before the id column existed have no id. For adding up a day's calories those are
    matched by display name (name_fallback=True): the worst case is a wrong number. For handing a file
    to someone (export) the name is NOT trusted, because anyone can set their display name to match
    another person's; there only the unforgeable id counts."""
    row_id = row[COL_UPLOADER_ID - 1] if len(row) >= COL_UPLOADER_ID else None
    if uploader_id and row_id not in (None, ""):
        return str(row_id).strip() == uploader_id
    return name_fallback and row[COL_UPLOADER - 1] == uploader


def _put_text(sheet: Worksheet, row: int, column: int, value: str) -> None:
    cell = sheet.cell(row=row, column=column)

    # Without the protection below, text starting with "=" is stored as a formula.
    cell.value = value

    # ==== 防護 4：寫進 Excel 的文字不被當成公式　開始 ====
    text = str(value)
    cell.value = text
    cell.data_type = "s"  # force "text", so Excel never evaluates it
    if text.startswith(FORMULA_STARTERS):
        cell.quotePrefix = True  # same as typing a leading ' in Excel: shows the text as-is
    # ==== 防護 4 結束 ====
